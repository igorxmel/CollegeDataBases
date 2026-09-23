"""Подготовка данных заказчика и загрузка их в базу данных «Мозаика».

Скрипт читает файлы с пометкой import из папки «Ресурсы», приводит данные
к виду, пригодному для загрузки (обрезка пробелов, округление цен до двух
знаков, перевод дат из формата Excel в ISO-8601), после чего:

  * создаёт файл базы данных mosaic.db по схеме из schema.sql;
  * сохраняет скрипт наполнения data.sql.

Запуск: python build_database.py
"""

from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import sqlite3

from openpyxl import load_workbook

DATABASE_DIR = Path(__file__).resolve().parent
RESOURCES_DIR = DATABASE_DIR.parents[1] / "Ресурсы"
SCHEMA_PATH = DATABASE_DIR / "schema.sql"
DATA_SCRIPT_PATH = DATABASE_DIR / "data.sql"
DATABASE_PATH = DATABASE_DIR / "mosaic.db"

# В Excel нумерация дней начинается с 1900 года, но с известной ошибкой
# високосного 1900 года, поэтому точкой отсчёта служит 30.12.1899.
EXCEL_EPOCH = date(1899, 12, 30)

# Справочники подсистем, для которых заказчик не передал файлов импорта.
# Значения взяты не произвольно, а из описания предметной области
# (Приложение 1): должности названы в п.2, п.5, п.8 и п.9, статусы заявки —
# в описании процесса работы с заявкой в п.3.
REFERENCE_BOOKS = {
    "job_position": (
        "job_position_name",
        ["Менеджер", "Мастер производства", "Аналитик", "Кладовщик"],
    ),
    "request_status": (
        "request_status_name",
        [
            "Создана",
            "Ожидает предоплаты",
            "В производстве",
            "Готова к отгрузке",
            "Выполнена",
            "Отменена",
        ],
    ),
}

# Порядок выгрузки в data.sql: родительские таблицы идут раньше дочерних,
# иначе внешние ключи не дадут выполнить скрипт наполнения.
TABLES_IN_LOAD_ORDER = [
    "unit",
    "material_type",
    "product_type",
    "organization_type",
    "job_position",
    "request_status",
    "supplier",
    "material",
    "material_supplier",
    "stock_movement",
    "employee",
    "product",
    "product_material",
    "production_batch",
    "partner",
    "partner_sale",
    "request",
    "request_item",
]


def read_rows(file_name):
    """Вернуть строки первого листа книги без строки заголовков."""
    workbook = load_workbook(RESOURCES_DIR / file_name, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    workbook.close()
    return [row for row in rows if any(cell is not None and str(cell).strip() for cell in row)]


def clean_text(value):
    return str(value).strip()


def to_money(value):
    """Округлить цену до двух знаков после запятой по правилам арифметики."""
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def to_iso_date(value):
    """Привести дату из ячейки Excel (число или дата) к строке YYYY-MM-DD.

    openpyxl возвращает ячейки с датой как datetime, поэтому время отбрасывается:
    в базе данных хранится только дата начала работы с поставщиком.
    """
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return (EXCEL_EPOCH + timedelta(days=int(value))).isoformat()


def collect_data():
    """Прочитать и подготовить данные всех импортируемых файлов."""
    material_types = [
        (clean_text(name), float(loss))
        for name, loss in read_rows("Material_type_import.xlsx")
    ]
    product_types = [
        (clean_text(name), float(coefficient))
        for name, coefficient in read_rows("Product_type_import.xlsx")
    ]
    suppliers = [
        (
            clean_text(name),
            clean_text(supplier_type),
            clean_text(inn),
            int(rating),
            to_iso_date(work_start_date),
        )
        for name, supplier_type, inn, rating, work_start_date in read_rows(
            "Suppliers_import.xlsx"
        )
    ]
    materials = [
        (
            clean_text(name),
            clean_text(material_type),
            to_money(price),
            float(stock),
            float(minimum),
            float(package),
            clean_text(unit),
            # Описание и изображение материала заказчик не предоставил файлом
            # импорта — в базе данных для них зарезервированы поля на будущее.
            None,
            None,
        )
        for name, material_type, price, stock, minimum, package, unit in read_rows(
            "Materials_import.xlsx"
        )
    ]
    material_suppliers = [
        (clean_text(material), clean_text(supplier))
        for material, supplier in read_rows("Material_suppliers_import.xlsx")
    ]

    # Справочники, которых нет отдельными файлами, собираются из данных.
    units = sorted({material[6] for material in materials})
    organization_types = sorted({supplier[1] for supplier in suppliers})

    return {
        "units": units,
        "material_types": material_types,
        "product_types": product_types,
        "organization_types": organization_types,
        "suppliers": suppliers,
        "materials": materials,
        "material_suppliers": material_suppliers,
    }


def check_references(data):
    """Проверить, что связи в исходных файлах не ссылаются на несуществующие записи."""
    material_type_names = {name for name, _ in data["material_types"]}
    material_names = {material[0] for material in data["materials"]}
    supplier_names = {supplier[0] for supplier in data["suppliers"]}

    problems = []
    for material in data["materials"]:
        if material[1] not in material_type_names:
            problems.append(f"материал «{material[0]}»: неизвестный тип «{material[1]}»")
    for material_name, supplier_name in data["material_suppliers"]:
        if material_name not in material_names:
            problems.append(f"связь: неизвестный материал «{material_name}»")
        if supplier_name not in supplier_names:
            problems.append(f"связь: неизвестный поставщик «{supplier_name}»")
    if problems:
        raise ValueError("Данные не прошли проверку:\n  " + "\n  ".join(problems))


def fill_database(connection, data):
    """Заполнить таблицы подготовленными данными."""
    cursor = connection.cursor()

    cursor.executemany("INSERT INTO unit (unit_name) VALUES (?)", [(u,) for u in data["units"]])
    cursor.executemany(
        "INSERT INTO material_type (material_type_name, loss_percent) VALUES (?, ?)",
        data["material_types"],
    )
    cursor.executemany(
        "INSERT INTO product_type (product_type_name, type_coefficient) VALUES (?, ?)",
        data["product_types"],
    )
    cursor.executemany(
        "INSERT INTO organization_type (organization_type_name) VALUES (?)",
        [(t,) for t in data["organization_types"]],
    )
    cursor.executemany(
        """
        INSERT INTO supplier (supplier_name, organization_type_id, inn, rating, work_start_date)
        VALUES (
            ?,
            (SELECT organization_type_id FROM organization_type
              WHERE organization_type_name = ?),
            ?, ?, ?
        )
        """,
        data["suppliers"],
    )
    cursor.executemany(
        """
        INSERT INTO material (
            material_name, material_type_id, unit_id,
            unit_price, stock_quantity, min_quantity, package_quantity,
            description, image_path
        )
        VALUES (
            ?,
            (SELECT material_type_id FROM material_type WHERE material_type_name = ?),
            (SELECT unit_id FROM unit WHERE unit_name = ?),
            ?, ?, ?, ?, ?, ?
        )
        """,
        [
            (name, material_type, unit, price, stock, minimum, package, description, image_path)
            for name, material_type, price, stock, minimum, package, unit, description, image_path
            in data["materials"]
        ],
    )
    cursor.executemany(
        """
        INSERT OR IGNORE INTO material_supplier (material_id, supplier_id)
        VALUES (
            (SELECT material_id FROM material WHERE material_name = ?),
            (SELECT supplier_id FROM supplier WHERE supplier_name = ?)
        )
        """,
        data["material_suppliers"],
    )
    connection.commit()


def reset_database(connection):
    """Удалить все таблицы базы данных перед созданием схемы заново.

    Файл базы может быть занят другой программой (например, открыт в клиенте
    СУБД), и тогда его не удалить средствами файловой системы. В этом случае
    схема пересоздаётся поверх существующего файла: на время очистки проверка
    внешних ключей отключается, иначе порядок удаления таблиц имел бы значение.
    """
    connection.execute("PRAGMA foreign_keys = OFF")
    tables = [
        row[0]
        for row in connection.execute(
            "SELECT name FROM sqlite_master WHERE type = 'table' "
            "AND name NOT LIKE 'sqlite_%'"
        )
    ]
    for table in tables:
        connection.execute(f'DROP TABLE IF EXISTS "{table}"')
    connection.commit()
    connection.execute("PRAGMA foreign_keys = ON")


def fill_reference_books(connection):
    """Заполнить справочники подсистем, описанные в предметной области.

    Остальные таблицы этих подсистем (продукция, партнёры, заявки, сотрудники)
    остаются пустыми: заказчик передал файлы импорта только по материалам,
    поставщикам и типам, а выдумывать за него данные нельзя.
    """
    for table, (column, values) in REFERENCE_BOOKS.items():
        connection.executemany(
            f"INSERT INTO {table} ({column}) VALUES (?)",
            [(value,) for value in values],
        )
    connection.commit()


def sql_literal(value):
    if value is None:
        return "NULL"
    if isinstance(value, str):
        return "'" + value.replace("'", "''") + "'"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def dump_data_script(connection):
    """Выгрузить содержимое таблиц в скрипт data.sql."""
    lines = [
        "-- ---------------------------------------------------------------------------",
        "-- Производственная компания «Мозаика»",
        "-- Данные заказчика, подготовленные из файлов с пометкой import.",
        "-- Скрипт выполняется после schema.sql.",
        "-- ---------------------------------------------------------------------------",
        "",
        "PRAGMA foreign_keys = ON;",
        "",
        "BEGIN TRANSACTION;",
        "",
    ]
    for table in TABLES_IN_LOAD_ORDER:
        cursor = connection.execute(f"SELECT * FROM {table}")
        columns = [description[0] for description in cursor.description]
        rows = cursor.fetchall()
        lines.append(f"-- Таблица {table}: {len(rows)} строк")
        for row in rows:
            values = ", ".join(sql_literal(value) for value in row)
            lines.append(f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({values});")
        lines.append("")
    lines.append("COMMIT;")
    lines.append("")
    DATA_SCRIPT_PATH.write_text("\n".join(lines), encoding="utf-8")


def main():
    data = collect_data()
    check_references(data)

    if DATABASE_PATH.exists():
        try:
            DATABASE_PATH.unlink()
        except PermissionError:
            # Файл открыт в другой программе: очистим его содержимое на месте.
            pass

    connection = sqlite3.connect(DATABASE_PATH)
    reset_database(connection)
    connection.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))
    fill_database(connection, data)
    fill_reference_books(connection)
    dump_data_script(connection)

    print(f"База данных создана: {DATABASE_PATH}")
    for table in TABLES_IN_LOAD_ORDER:
        count = connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        print(f"  {table}: {count}")
    connection.close()
    print(f"Скрипт наполнения сохранён: {DATA_SCRIPT_PATH}")


if __name__ == "__main__":
    main()

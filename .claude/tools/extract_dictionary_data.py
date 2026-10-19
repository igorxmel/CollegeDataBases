"""One-off: rebuild `dictionary_data.py` from the delivered Task 3 report.

The report was regenerated elsewhere against the new 18-table schema while the local
generator still described the old 7-table one. Rather than retyping everything, the data
is read back out of the finished document, so the generator reproduces exactly what was
handed in.

Usage:
    python .claude/tools/extract_dictionary_data.py
"""

import io
import re
import sys
from pathlib import Path

from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "Task 3rd Database Documentation" / "Solution" / "Отчёт по ПР №3.docx"
OUT = Path(__file__).resolve().parent / "dictionary_data.py"

doc = Document(REPORT)

# --- таблицы документа по порядку
tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in doc.tables]

# первая таблица это связи, дальше идут описания полей
relations = [tuple(r) for r in tables[0][1:]]

# --- проза: суть проекта и вывод
paras = [p.text.strip() for p in doc.paragraphs]


def section(start: str, stop: str) -> list[str]:
    out, on = [], False
    for text in paras:
        if text.startswith(start):
            on = True
            continue
        if on and text.startswith(stop):
            break
        if on and text and not text.startswith(("Таблица ", "Продолжение таблицы")):
            out.append(text)
    return out


intro = section("1 Суть проекта", "2 Описание связей")
conclusion = section("4 Вывод", "\u0000")

# --- описания таблиц: имя, русское название, комментарий, поля
CAPTION = re.compile(r"^Таблица (\d+) - Описание таблицы (\S+) - (.+)$")
LEAD = re.compile(r"^Таблица (\S+) - (.+?)\.$")
COMMENT = re.compile(r"^Комментарий к таблице: (.+)$")

order: list[str] = []
meta: dict[str, dict] = {}
current = None
table_index = 1  # tables[0] уже израсходована под связи

for text in paras:
    m = CAPTION.match(text)
    if m:
        _, name, rus = m.groups()
        current = name
        entry = meta.setdefault(name, {"rus": rus, "comment": "", "fields": []})
        if name not in order:
            order.append(name)
        # столбцы таблицы описания, первый столбец это боковик
        grid = tables[table_index]
        table_index += 1
        for col in range(1, len(grid[0])):
            entry["fields"].append(tuple(row[col] for row in grid))
        continue
    if text.startswith("Продолжение таблицы") and current:
        grid = tables[table_index]
        table_index += 1
        for col in range(1, len(grid[0])):
            meta[current]["fields"].append(tuple(row[col] for row in grid))
        continue
    m = COMMENT.match(text)
    if m and current:
        meta[current]["comment"] = m.group(1)

# --- запись модуля
def lit(value) -> str:
    return repr(value).replace("'", '"') if isinstance(value, str) else repr(value)


lines = [
    '"""Данные словаря данных, вынуты из сданного отчёта по ПР №3.',
    "",
    "Файл собран скриптом extract_dictionary_data.py и правится руками только вместе",
    "с отчётом: build_dictionary.py собирает документ именно отсюда.",
    '"""',
    "",
    "INTRO = [",
]
lines += [f"    {lit(p)}," for p in intro]
lines += ["]", "", "RELATIONS = ["]
lines += [f"    {r!r}," for r in relations]
lines += ["]", "", "# (имя, русское название, комментарий, поля)", "TABLES = ["]
for name in order:
    e = meta[name]
    lines.append(f"    ({lit(name)}, {lit(e['rus'])}, {lit(e['comment'])}, [")
    for field in e["fields"]:
        lines.append(f"        {field!r},")
    lines.append("    ]),")
lines += ["]", "", "CONCLUSION = ["]
lines += [f"    {lit(p)}," for p in conclusion]
lines += ["]", ""]

OUT.write_text("\n".join(lines), encoding="utf-8")

print(f"таблиц описано: {len(order)}")
print(f"связей: {len(relations)}")
print(f"полей всего: {sum(len(meta[n]['fields']) for n in order)}")
print(f"абзацев вступления: {len(intro)}, вывода: {len(conclusion)}")
print(f"записано: {OUT.name}")

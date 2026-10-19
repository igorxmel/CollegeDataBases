"""Build the Task 5 report (parts 5.1 and 5.2) as a .docx file.

The SQL is not duplicated here: the script reads `Solution/5.1_create.sql` and
`Solution/5.2_modify.sql` and embeds them, so the report can never drift from the
scripts that actually run.

Usage:
    python .claude/tools/build_sql_report.py
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TASK = Path(__file__).resolve().parents[2] / "Task 5th MySQL Create and Modify"
SOLUTION = TASK / "Solution"
OUT = SOLUTION / "Отчёт по ПР №5.docx"

FONT = "Times New Roman"
MONO = "Courier New"
SIZE = Pt(14)
CODE_SIZE = Pt(9)
TABLE_SIZE = Pt(10)
INDENT = Cm(1.25)
TEXT_WIDTH = 16.5

TITLE = [
    "Министерство общего и профессионального образования",
    "Ростовской области",
    "Государственное бюджетное профессиональное образовательное",
    "учреждение Ростовской области",
    "«Ростовский-на-Дону колледж связи и информатики»",
]
SIGN_LEFT = ["Выполнил: студент 4 курса", "группы ИС-42", "Холкин Егор"]
SIGN_RIGHT = ["Проверил: преподаватель", "Кротенко Е. М."]

GOAL = (
    "Научиться создавать базу данных и таблицы в MySQL средствами языка SQL, задавать "
    "ограничения целостности, а затем изменять структуру готовых таблиц командой "
    "ALTER TABLE и автоматизировать заполнение полей триггерами."
)

TASK_TEXT = [
    "Часть 5.1: создать базу данных «torgash», sql-код сохранить в отчёт. По условию имя "
    "базы дополняется фамилией, поэтому база называется torgash_Холкин. Требуется шесть "
    "таблиц: виды собственности, продавцы, заказчики физические лица, заказчики "
    "юридические лица, заказы и скидки, с ограничениями, перечисленными в задании.",

    "Часть 5.2: законспектировать команды изменения структуры таблицы и выполнить "
    "семь изменений над созданной базой: уникальность названия продавца, дата "
    "регистрации, значение по умолчанию для формы собственности, возраст заказчика, "
    "сообщение о пенсионном возрасте, статус юридического лица и его автоматическое "
    "выставление по рейтингу.",
]

COMMANDS = [
    ("ALTER TABLE <таблица> ADD <столбец> <тип>",
     "Добавление нового столбца. FIRST добавляет столбец первым, AFTER <столбец> после "
     "указанного, без уточнения столбец идёт в конец таблицы"),
    ("ALTER TABLE <таблица> MODIFY COLUMN <столбец> <тип> [атрибуты]",
     "Изменение типа столбца и его атрибутов"),
    ("ALTER TABLE <таблица> ALTER COLUMN <столбец> SET DEFAULT <значение>",
     "Изменение значения по умолчанию"),
    ("ALTER TABLE <таблица> CHANGE <старое имя> <новое имя> <тип>",
     "Переименование столбца, тип указывать обязательно"),
    ("ALTER TABLE <таблица> DROP <столбец>",
     "Удаление столбца"),
    ("ALTER TABLE <таблица> ADD INDEX [имя] (столбец)",
     "Добавление индекса, без имени индекс называется по имени столбца"),
    ("ALTER TABLE <таблица> DROP INDEX <имя>",
     "Удаление индекса"),
    ("ALTER TABLE <таблица> ADD PRIMARY KEY (столбец)",
     "Назначение первичного ключа"),
    ("ALTER TABLE <таблица> DROP PRIMARY KEY",
     "Удаление первичного ключа"),
    ("ALTER TABLE <таблица> ADD CONSTRAINT <имя> FOREIGN KEY (столбец) "
     "REFERENCES <таблица> (столбец)",
     "Добавление внешнего ключа с именем ограничения"),
    ("ALTER TABLE <таблица> DROP FOREIGN KEY <имя>",
     "Удаление внешнего ключа по имени ограничения"),
    ("ALTER TABLE <старое имя> RENAME <новое имя>",
     "Переименование таблицы"),
    ("DROP TABLE <таблица>",
     "Удаление таблицы"),
    ("DELETE FROM <таблица> [WHERE <условие>]",
     "Удаление строк: без условия удаляются все, с условием только подходящие"),
    ("UPDATE <таблица> SET <столбец> = <значение> WHERE <условие>",
     "Изменение содержимого таблицы"),
]

NOTES_51 = [
    "Порядок создания таблиц задают внешние ключи: сначала справочники vids и skid, "
    "затем продавцы и заказчики, последними заказы.",

    "В справочник скидок я добавил столбец summa с числовым порогом. В задании условие "
    "записано текстом («Общая стоимость заказов больше 1500 руб»), а триггеру нужно "
    "число, и разбирать строку при каждом заказе дороже, чем хранить порог рядом.",

    "Форма собственности вынесена в отдельный столбец id_vs со ссылкой на справочник, "
    "поэтому в названиях продавцов и юридических лиц приставки ИП, ООО и ПАО убраны.",

    "Возраст не меньше 18 лет проверяется триггером, а не ограничением CHECK: MySQL не "
    "разрешает использовать внутри CHECK недетерминированные функции, к которым "
    "относится CURDATE().",

    "Запреты на изменение (комиссионные, пол, дата рождения) сделаны триггерами "
    "BEFORE UPDATE с вызовом SIGNAL SQLSTATE '45000', который прерывает операцию с "
    "сообщением.",

    "Скидка физического лица пересчитывается триггером AFTER INSERT на таблице заказов: "
    "берётся максимальная скидка справочника, порог которой не превышает сумму всех "
    "заказов этого заказчика.",
]

NOTES_52 = [
    "Триггеры, созданные в части 5.1, здесь пересоздаются через DROP TRIGGER и CREATE "
    "TRIGGER: изменить тело существующего триггера MySQL не умеет.",

    "Сообщение о пенсионном возрасте выдаётся через SIGNAL SQLSTATE '01000'. Это класс "
    "предупреждения, поэтому строка вставляется, а текст виден по команде SHOW WARNINGS. "
    "Код '45000' в этом месте прервал бы вставку, что заданием не требуется.",

    "Столбец data_reg объявлен с DEFAULT (CURRENT_DATE), такая запись значения по "
    "умолчанию выражением поддерживается начиная с MySQL 8.0.13.",

    "Столбец vozrast заполняется триггером при вставке, а для уже внесённых заказчиков "
    "пересчитан командой UPDATE. Вычисляемым столбцом его сделать нельзя по той же "
    "причине, что и с CHECK: CURDATE() недетерминированная.",
]

CONCLUSION = [
    "В работе я создал базу данных torgash_Холкин из шести таблиц и наполнил её данными "
    "из задания, а затем изменил структуру трёх таблиц командой ALTER TABLE и "
    "автоматизировал заполнение полей триггерами.",

    "Главное, что я вынес: часть требований задания на уровне ограничений CHECK не "
    "выражается. Проверка возраста, автоматическая скидка, запрет на изменение уже "
    "внесённых значений и бесплатная доставка для дорогих заказов это всё триггеры, "
    "потому что ограничение проверяет значение само по себе, а здесь нужно сравнение с "
    "текущей датой, с содержимым другой таблицы или со старым значением строки.",

    "Отдельно отмечу разницу между SIGNAL с кодом '45000' и '01000': первый прерывает "
    "операцию, второй только выводит предупреждение. В задании про пенсионный возраст "
    "требуется именно сообщение, а не отказ, поэтому там второй вариант.",
]


def setup(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(doc, text="", *, indent=INDENT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = indent
    p.paragraph_format.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = SIZE
        run.bold = bold
    return p


def code(doc, text: str) -> None:
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line.strip() else "")
        run.font.name = MONO
        run.font.size = CODE_SIZE
        run.element.rPr.rFonts.set(qn("w:eastAsia"), MONO)


def blank(doc, n=1):
    for _ in range(n):
        para(doc)


def heading(doc, number, title, *, new_page=True):
    if new_page:
        para(doc).add_run().add_break(WD_BREAK.PAGE)
    para(doc, f"{number} {title}", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)


def sub(doc, number, title):
    para(doc, f"{number} {title}", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)


def build():
    doc = Document()
    setup(doc)

    for line in TITLE:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    blank(doc, 3)
    para(doc, "ОТЧЁТ", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    para(doc, "по практической работе № 5", indent=Cm(0),
         align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    blank(doc)
    para(doc, "Тема: Создание и модификация базы данных в MySQL", indent=Cm(0),
         align=WD_ALIGN_PARAGRAPH.CENTER)
    blank(doc, 4)
    for line in SIGN_LEFT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    blank(doc)
    for line in SIGN_RIGHT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    blank(doc, 2)
    para(doc, "Ростов-на-Дону", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "2026", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)

    heading(doc, "1", "Цель работы")
    para(doc, GOAL)

    heading(doc, "2", "Задание")
    for text in TASK_TEXT:
        para(doc, text)

    heading(doc, "3", "Конспект команд изменения структуры таблицы")
    para(doc, "Команды, разобранные в практической работе, сведены в таблице 1.")
    blank(doc)
    para(doc, "Таблица 1 - Команды изменения структуры таблицы", indent=Cm(0))

    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for i, (syntax, description) in enumerate([("Синтаксис", "Описание")] + COMMANDS):
        cells = t.add_row().cells
        for cell, value, width, mono in (
            (cells[0], syntax, 8.5, i > 0),
            (cells[1], description, 8.0, False),
        ):
            cell.width = Cm(width)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            run.font.name = MONO if mono else FONT
            run.font.size = CODE_SIZE if mono else TABLE_SIZE
            run.bold = i == 0
    blank(doc)

    heading(doc, "4", "Часть 5.1. Создание базы данных")
    sub(doc, "4.1", "Пояснения к решению")
    for text in NOTES_51:
        para(doc, text)
    blank(doc)
    sub(doc, "4.2", "SQL-код")
    code(doc, (SOLUTION / "5.1_create.sql").read_text(encoding="utf-8"))

    heading(doc, "5", "Часть 5.2. Модификация базы данных")
    sub(doc, "5.1", "Пояснения к решению")
    for text in NOTES_52:
        para(doc, text)
    blank(doc)
    sub(doc, "5.2", "SQL-код")
    code(doc, (SOLUTION / "5.2_modify.sql").read_text(encoding="utf-8"))

    heading(doc, "6", "Вывод")
    for text in CONCLUSION:
        para(doc, text)

    try:
        doc.save(OUT)
    except PermissionError:
        raise SystemExit(f"{OUT.name} открыт в Word, закрой его и повтори")
    print(f"saved: {OUT.name}")


if __name__ == "__main__":
    build()

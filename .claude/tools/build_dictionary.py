"""Build the Task 3 report: словарь данных for the «Мозаика» database from Task 2.

Structure follows the template in the assignment: краткое описание проекта, описание
связей одной таблицей, затем по таблице на каждую таблицу базы, где строки это
характеристики поля, а столбцы сами поля.

Текст и поля лежат в `dictionary_data.py` рядом: русских названий и комментариев в
`schema.sql` нет, поэтому они хранятся отдельно. Тот файл собирается скриптом
`extract_dictionary_data.py` из сданного отчёта, руками его править не нужно.

Usage:
    python .claude/tools/build_dictionary.py
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from dictionary_data import CONCLUSION, INTRO, RELATIONS, TABLES, TABLES_LEAD

OUT = (
    Path(__file__).resolve().parents[2]
    / "Task 3rd Database Documentation"
    / "Solution"
    / "Отчёт по ПР №3.docx"
)

FONT = "Times New Roman"
SIZE = Pt(14)
TABLE_SIZE = Pt(10)
INDENT = Cm(1.25)
TEXT_WIDTH = 16.5
MAX_FIELDS = 5  # столбцов-полей в одной части таблицы

TITLE = [
    "Министерство общего и профессионального образования",
    "Ростовской области",
    "Государственное бюджетное профессиональное образовательное",
    "учреждение Ростовской области",
    "«Ростовский-на-Дону колледж связи и информатики»",
]
SIGN_LEFT = ["Выполнил: студент 4 курса", "группы ИС-42", "Холкин Егор"]
SIGN_RIGHT = ["Проверил: преподаватель", "Кротенко Е. М."]



# поле: (имя в ER, русское название, тип, размерность, обязательность,
#        ограничение, значение по умолчанию, тип ключа, комментарий)

ROW_LABELS = [
    "Название поля в ER-модели",
    "Русскоязычное название поля",
    "Тип данных поля",
    "Размерность поля",
    "Обязательность поля",
    "Ограничение на ввод данных в поле",
    "Значение по умолчанию",
    "Тип ключа",
    "Комментарий к полю",
]



def setup(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(doc, text="", *, indent=INDENT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False,
         size=SIZE):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = indent
    p.paragraph_format.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = size
        run.bold = bold
    return p


def blank(doc, n=1):
    for _ in range(n):
        para(doc)


def heading(doc, number, title, *, new_page=True):
    if new_page:
        para(doc).add_run().add_break(WD_BREAK.PAGE)
    para(doc, f"{number} {title}", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)


def grid(doc, rows, widths, *, header=False):
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for i, values in enumerate(rows):
        cells = t.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Cm(width)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.name = FONT
            run.font.size = TABLE_SIZE
            run.bold = header and i == 0
    blank(doc)


def field_table(doc, number, name, rus, fields, part=None):
    """Одна часть таблицы описания: строки это характеристики, столбцы это поля."""
    if part:
        label = f"Продолжение таблицы {number}"
    else:
        label = f"Таблица {number} - Описание таблицы {name} - {rus}"
    para(doc, label, indent=Cm(0))

    label_width = 4.6
    field_width = (TEXT_WIDTH - label_width) / len(fields)
    rows = [
        [ROW_LABELS[i]] + [field[i] for field in fields]
        for i in range(len(ROW_LABELS))
    ]
    grid(doc, rows, [label_width] + [field_width] * len(fields))


def build():
    doc = Document()
    setup(doc)

    for line in TITLE:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    blank(doc, 3)
    para(doc, "ОТЧЁТ", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    para(doc, "по практической работе № 3", indent=Cm(0),
         align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    blank(doc)
    para(doc, "Тема: Составление технической документации проекта", indent=Cm(0),
         align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "«Вариант 1. Подсистема учёта материалов компании «Мозаика»»",
         indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    blank(doc, 4)
    for line in SIGN_LEFT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    blank(doc)
    for line in SIGN_RIGHT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    blank(doc, 2)
    para(doc, "Ростов-на-Дону", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "2026", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)

    heading(doc, "1", "Суть проекта")
    for text in INTRO:
        para(doc, text)

    heading(doc, "2", "Описание связей")
    para(doc, "Связи между таблицами базы данных приведены в таблице 1.")
    blank(doc)
    para(doc, "Таблица 1 - Описание связей между таблицами", indent=Cm(0))
    grid(
        doc,
        [("Пара таблиц", "Тип отношения", "Пояснение", "Каскадная операция")] + RELATIONS,
        [4.2, 2.6, 4.6, 5.1],
        header=True,
    )

    heading(doc, "3", "Описание таблиц")
    para(doc, TABLES_LEAD)
    para(doc,
         f"Описание таблиц приведено в таблицах со 2 по {len(TABLES) + 1}. Таблицы, в "
         f"которых полей больше пяти, разбиты на части: продолжение помечено словами "
         f"«Продолжение таблицы» и повторяет боковик.")
    blank(doc)

    number = 2
    for name, rus, comment, fields in TABLES:
        para(doc, f"Таблица {name} - {rus}.")
        blank(doc)
        parts = [fields[i:i + MAX_FIELDS] for i in range(0, len(fields), MAX_FIELDS)]
        for index, chunk in enumerate(parts, start=1):
            part = index > 1
            field_table(doc, number, name, rus, chunk, part)
        para(doc, f"Комментарий к таблице: {comment}")
        blank(doc)
        number += 1

    heading(doc, "4", "Вывод")
    for text in CONCLUSION:
        para(doc, text)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(OUT)
    except PermissionError:
        raise SystemExit(f"{OUT.name} открыт в Word, закрой его и повтори")
    print(f"saved: {OUT.name}")


if __name__ == "__main__":
    build()
